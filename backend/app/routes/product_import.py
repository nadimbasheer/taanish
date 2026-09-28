from io import BytesIO
import re
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.product import Product
from app.models.inventory import StockMovement
from app.services.dependencies import require_role


router = APIRouter(
    prefix="/products",
    tags=["Product Import"]
)


def clean_value(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def decimal_value(value, field_name):
    value = clean_value(value)

    if value is None:
        raise HTTPException(
            status_code=400,
            detail=f"{field_name} is required"
        )

    try:
        number = Decimal(str(value))

        if number < 0:
            raise HTTPException(
                status_code=400,
                detail=f"{field_name} cannot be negative"
            )

        return number

    except InvalidOperation:
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {field_name}: {value}"
        )


def integer_value(value, field_name):
    value = clean_value(value)

    if value is None:
        raise HTTPException(
            status_code=400,
            detail=f"{field_name} is required"
        )

    try:
        number = int(value)

        if number < 0:
            raise HTTPException(
                status_code=400,
                detail=f"{field_name} cannot be negative"
            )

        return number

    except (ValueError, TypeError):
        raise HTTPException(
            status_code=400,
            detail=f"Invalid {field_name}: {value}"
        )


def generate_next_barcode(db: Session):
    products = db.query(Product.barcode).all()

    highest = 0

    for row in products:
        barcode = row[0]

        if not barcode:
            continue

        match = re.fullmatch(r"FS(\d{6})", str(barcode))

        if match:
            number = int(match.group(1))
            highest = max(highest, number)

    next_number = highest + 1

    while True:
        barcode = f"FS{next_number:06d}"

        exists = (
            db.query(Product)
            .filter(Product.barcode == barcode)
            .first()
        )

        if not exists:
            return barcode

        next_number += 1


@router.post("/import-excel")
async def import_products_from_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user=Depends(require_role("admin", "manager"))
):
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Please select an Excel file"
        )

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx Excel files are supported"
        )

    try:
        contents = await file.read()

        workbook = load_workbook(
            filename=BytesIO(contents),
            data_only=True
        )

        sheet = workbook.active

        headers = {}

        for cell in sheet[1]:
            if cell.value is not None:
                headers[str(cell.value).strip().lower()] = cell.column

        required_columns = [
            "name",
            "selling",
            "stock",
            "min_stock"
        ]

        missing = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Missing columns: "
                    + ", ".join(missing)
                )
            )

        rows = []

        for row_number in range(2, sheet.max_row + 1):

            name = clean_value(
                sheet.cell(
                    row=row_number,
                    column=headers["name"]
                ).value
            )

            if not name:
                continue

            category = None
            sku = None
            barcode = None

            if "category" in headers:
                category = clean_value(
                    sheet.cell(
                        row=row_number,
                        column=headers["category"]
                    ).value
                )

            if "sku" in headers:
                sku = clean_value(
                    sheet.cell(
                        row=row_number,
                        column=headers["sku"]
                    ).value
                )

            if "barcode" in headers:
                barcode = clean_value(
                    sheet.cell(
                        row=row_number,
                        column=headers["barcode"]
                    ).value
                )

            cost = Decimal("0")

            if "cost" in headers:
                cost = decimal_value(
                    sheet.cell(
                        row=row_number,
                        column=headers["cost"]
                    ).value,
                    "cost"
                )

            mrp = Decimal("0")

            if "mrp" in headers:
                mrp = decimal_value(
                    sheet.cell(
                        row=row_number,
                        column=headers["mrp"]
                    ).value,
                    "mrp"
                )

            selling = decimal_value(
                sheet.cell(
                    row=row_number,
                    column=headers["selling"]
                ).value,
                "selling"
            )

            stock = integer_value(
                sheet.cell(
                    row=row_number,
                    column=headers["stock"]
                ).value,
                "stock"
            )

            min_stock = integer_value(
                sheet.cell(
                    row=row_number,
                    column=headers["min_stock"]
                ).value,
                "min_stock"
            )

            rows.append({
                "row": row_number,
                "name": str(name),
                "category": category,
                "sku": str(sku) if sku is not None else None,
                "barcode": str(barcode) if barcode else None,
                "cost": cost,
                "mrp": mrp,
                "selling": selling,
                "stock": stock,
                "min_stock": min_stock
            })

        if not rows:
            raise HTTPException(
                status_code=400,
                detail="No products found in the Excel file"
            )

        # Check duplicate SKUs inside Excel
        skus = [
            row["sku"]
            for row in rows
            if row["sku"]
        ]

        if len(skus) != len(set(skus)):
            raise HTTPException(
                status_code=400,
                detail="Duplicate SKU found inside Excel file"
            )

        # Check duplicate barcodes inside Excel
        barcodes = [
            row["barcode"]
            for row in rows
            if row["barcode"]
        ]

        if len(barcodes) != len(set(barcodes)):
            raise HTTPException(
                status_code=400,
                detail="Duplicate barcode found inside Excel file"
            )

        # Check SKUs against database
        for row in rows:
            if row["sku"]:
                existing = (
                    db.query(Product)
                    .filter(Product.sku == row["sku"])
                    .first()
                )

                if existing:
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"SKU already exists: "
                            f"{row['sku']}"
                        )
                    )

        # Check manually supplied barcodes against database
        for row in rows:
            if row["barcode"]:
                existing = (
                    db.query(Product)
                    .filter(
                        Product.barcode == row["barcode"]
                    )
                    .first()
                )

                if existing:
                    raise HTTPException(
                        status_code=400,
                        detail=(
                            f"Barcode already exists: "
                            f"{row['barcode']}"
                        )
                    )

        created_products = []

        try:
            for row in rows:

                barcode = row["barcode"]

                if not barcode:
                    barcode = generate_next_barcode(db)

                product = Product(
                    name=row["name"],
                    category=row["category"],
                    sku=row["sku"],
                    barcode=barcode,
                    cost=row["cost"],
                    mrp=row["mrp"],
                    selling=row["selling"],
                    stock=row["stock"],
                    min_stock=row["min_stock"]
                )

                db.add(product)
                db.flush()

                if row["stock"] > 0:
                    movement = StockMovement(
                        product_id=product.id,
                        movement_type="ADJUSTMENT",
                        quantity=row["stock"],
                        stock_before=0,
                        stock_after=row["stock"],
                        note="Opening Stock - Excel Import"
                    )

                    db.add(movement)

                created_products.append({
                    "id": product.id,
                    "name": product.name,
                    "sku": product.sku,
                    "barcode": product.barcode,
                    "stock": product.stock
                })

            db.commit()

        except Exception:
            db.rollback()
            raise

        return {
            "message": "Products imported successfully",
            "total_imported": len(created_products),
            "products": created_products
        }

    except HTTPException:
        raise

    except Exception as error:
        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=f"Import failed: {str(error)}"
        )